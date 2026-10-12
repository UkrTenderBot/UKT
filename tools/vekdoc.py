from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DIRECTOR = 'Ірина ПАЦАЙ'

def new():
    d = Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.5), Cm(1.5), Cm(1.5), Cm(1.5)
    st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    return d

def p(d, text='', bold=False, align=None, indent=True, size=None, italic=False, space=0, left=None):
    para = d.add_paragraph(); para.paragraph_format.space_after = Pt(space)
    if align is not None: para.alignment = align
    elif indent:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; para.paragraph_format.first_line_indent = Cm(1)
    if left is not None: para.paragraph_format.left_indent = Cm(left)
    r = para.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return para

def header(d):
    R = WD_ALIGN_PARAGRAPH.RIGHT
    p(d, 'ТОВАРИСТВО З ОБМЕЖЕНОЮ ВІДПОВІДАЛЬНІСТЮ «ЕЛЕКТРОТЕХНІЧНА КОМПАНІЯ «ВЕЛЕС»»', bold=True, align=R)
    p(d, '01103, м. Київ, Залізничне шосе, буд. 45, оф. 33-А; код ЄДРПОУ 43349778; ІПН 433497726551', italic=True, align=R, size=10)
    p(d, 'IBAN UA593348510000000026006152053 у ПАТ «ПУМБ», МФО 334851', italic=True, align=R, size=10)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = 'e-mail: patriot-cable@ukr.net'; t.rows[0].cells[1].text = 'тел. (050) 949-79-38'
    t.rows[0].cells[1].paragraphs[0].alignment = R
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs: r.italic = True; r.font.size = Pt(10)
    p(d, '_' * 78, align=WD_ALIGN_PARAGRAPH.CENTER, space=8)

def signature(d):
    tbl = d.add_table(rows=3, cols=3)
    vals = [['Директор', '______________', DIRECTOR], ['', '(підпис)', ''], ['', 'М.П.', '']]
    al = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            cell.text = vals[ri][ci]
            para = cell.paragraphs[0]; para.alignment = al[ci]
            for r in para.runs: r.font.size = Pt(12 if ri == 0 or ci == 1 and ri == 2 else 9)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def table(d, head, rows, widths=None, size=10):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = h; shade(c, 'D9E2F3')
        for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for r in cells[i].paragraphs[0].runs: r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(4)
    return t
